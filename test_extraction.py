import unittest
from io import BytesIO
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook
from PIL import Image
from streamlit.testing.v1 import AppTest

from extraction import Plate, Well, export_workbook, export_xlsx, extract_plate, normalize_image, validate_plate


ROWS = list("ABCDEFGH")
COLUMNS = [str(column) for column in range(1, 13)]


class PlateTests(unittest.TestCase):
    def test_multiple_worksheets(self):
        sheets = [
            ("photo/" + "a" * 40, [[value]], set(), ["AA"], ["24"])
            for value in (0.1, 0.2)
        ]
        workbook = load_workbook(BytesIO(export_workbook(sheets)))
        self.assertEqual(len(workbook.sheetnames), 2)
        self.assertTrue(all(len(name) <= 31 and "/" not in name for name in workbook.sheetnames))
        self.assertEqual([sheet["B2"].value for sheet in workbook], [0.1, 0.2])

    def plate(self):
        return Plate(rows=ROWS, columns=COLUMNS, wells=[
            Well(row=row, column=column, value=0.165 if (row, column) == ("A", "1") else None,
                 status="readable" if (row, column) == ("A", "1") else "empty")
            for row in ROWS for column in COLUMNS
        ])

    def test_mapping_and_uncertainty(self):
        plate = validate_plate(self.plate())
        values = [[None] * 12 for _ in ROWS]
        for well in plate.wells:
            values[ROWS.index(well.row)][COLUMNS.index(well.column)] = well.value
        values[3][11] = 0.033
        workbook = load_workbook(BytesIO(export_xlsx(values, {("B", "2")}, plate.rows, plate.columns)))
        sheet = workbook.active
        self.assertEqual(sheet["B2"].value, 0.165)
        self.assertEqual(sheet["M5"].value, 0.033)
        self.assertEqual(sheet["M1"].value, "12")
        self.assertEqual(sheet["A9"].value, "H")
        self.assertIsNone(sheet["C3"].value)
        self.assertIsNotNone(sheet["C3"].comment)

    def test_reject_bad_coordinates_and_guessed_values(self):
        plate = self.plate()
        plate.wells[-1] = plate.wells[0]
        with self.assertRaises(ValueError):
            validate_plate(plate)

    def test_larger_grid_preserves_labels_and_order(self):
        rows = list("ABCDEFGHIJKLMNOP") + ["AA", "AB"]
        columns = [str(column) for column in range(1, 25)]
        plate = validate_plate(Plate(rows=rows, columns=columns, wells=[
            Well(row=row, column=column, value=0.125, status="readable")
            for row in reversed(rows) for column in reversed(columns)
        ]))
        sheet = load_workbook(BytesIO(export_xlsx(
            [[0.125] * 24 for _ in rows], set(), plate.rows, plate.columns,
        ))).active
        self.assertEqual(sheet["A10"].value, "I")
        self.assertEqual(sheet["A19"].value, "AB")
        self.assertEqual(sheet["Y1"].value, "24")
        self.assertEqual(sheet["Y19"].value, 0.125)
        plate.rows.append("AB")
        with self.assertRaises(ValueError):
            validate_plate(plate)
        plate = self.plate()
        plate.wells[0].status = "unclear"
        with self.assertRaises(ValueError):
            validate_plate(plate)

    def test_image_and_api_contract(self):
        data = BytesIO()
        Image.new("RGB", (20, 20)).save(data, format="JPEG")
        image = normalize_image(data.getvalue())
        self.assertTrue(image.startswith(b"\x89PNG"))
        with self.assertRaises(ValueError):
            normalize_image(b"not an image")
        with patch("extraction.genai.Client") as client:
            api = client.return_value.__enter__.return_value
            api.models.generate_content.return_value = MagicMock(text=self.plate().model_dump_json())
            self.assertEqual(extract_plate(image, "test-key", "test-model").wells[0].value, 0.165)
            args = api.models.generate_content.call_args.kwargs
            self.assertTrue(args["config"].automatic_function_calling.disable)
            self.assertEqual(args["config"].response_mime_type, "application/json")
            self.assertEqual(args["contents"][1].inline_data.mime_type, "image/png")
            api.models.generate_content.return_value.text = None
            with self.assertRaises(ValueError):
                extract_plate(image, "test-key", "test-model")

    def test_ui_upload_and_camera_modes(self):
        with patch.dict("os.environ", {"GEMINI_API_KEY": ""}):
            app = AppTest.from_file("app.py", default_timeout=15).run()
            self.assertFalse(app.exception)
            self.assertEqual(app.title[0].value, "Plate to Excel")
            self.assertEqual(app.radio[0].value, "Take photo")
            self.assertTrue(app.button[0].disabled)
            self.assertFalse(app.exception)
            self.assertTrue(any("camera access" in item.value for item in app.caption))

    def test_ui_extract_and_download(self):
        photo = BytesIO()
        Image.new("RGB", (20, 20)).save(photo, format="JPEG")
        photo.name = "plate.jpg"
        with patch.dict("os.environ", {"GEMINI_API_KEY": "test-key"}), patch(
            "streamlit.file_uploader", return_value=[photo, photo]
        ), patch("extraction.extract_plate", return_value=self.plate()) as extract:
            app = AppTest.from_file("app.py", default_timeout=15).run()
            app.radio[0].set_value("Upload photos").run()
            app.button[0].click().run()
            self.assertFalse(app.exception)
            self.assertEqual(extract.call_count, 2)
            self.assertEqual(len(app.dataframe), 2)
            self.assertEqual(app.get("download_button")[0].label, "Download Excel")
            self.assertEqual(app.dataframe[0].value.iloc[0, 0], 0.165)

    def test_camera_queue_extract_remove_and_clear(self):
        photo = BytesIO()
        Image.new("RGB", (20, 20)).save(photo, format="JPEG")
        with patch.dict("os.environ", {"GEMINI_API_KEY": "test-key"}), patch(
            "streamlit.camera_input", return_value=photo
        ) as camera, patch("extraction.extract_plate", return_value=self.plate()) as extract:
            app = AppTest.from_file("app.py", default_timeout=15).run()
            for _ in range(2):
                next(button for button in app.button if button.label == "Add photo").click().run()
            self.assertFalse(app.exception)
            self.assertEqual(len(app.session_state.captures), 2)
            self.assertEqual(camera.call_args.kwargs["key"], "camera-3")
            app.run()
            self.assertEqual(len(app.session_state.captures), 2)
            app.radio[0].set_value("Upload photos").run()
            app.radio[0].set_value("Take photo").run()
            self.assertEqual(len(app.session_state.captures), 2)
            next(button for button in app.button if button.label == "Extract values").click().run()
            self.assertFalse(app.exception)
            self.assertEqual(extract.call_count, 2)
            self.assertEqual(len(app.dataframe), 2)
            next(button for button in app.button if button.label == "Remove photo").click().run()
            self.assertEqual(len(app.session_state.captures), 1)
            self.assertEqual(len(app.get("download_button")), 0)
            next(button for button in app.button if button.label == "Clear captured photos").click().run()
            self.assertFalse(app.exception)
            self.assertEqual(len(app.session_state.captures), 0)


if __name__ == "__main__":
    unittest.main()
