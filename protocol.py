from collections import defaultdict
from datetime import datetime
from io import BytesIO
from statistics import mean, stdev
from typing import Literal

from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from pydantic import BaseModel, Field, model_validator

from extraction import Plate, write_sheet


class ReadingTemplate(BaseModel):
    name: str = Field(min_length=1, max_length=80)
    flask_start: int = Field(default=1, ge=1, le=10000)
    flask_count: int = Field(default=8, ge=1, le=384)
    replicates: int = Field(default=2, ge=1, le=48)
    row_direction: Literal['top', 'bottom'] = 'bottom'
    column_direction: Literal['left', 'right'] = 'left'
    start_slot: int = Field(default=0, ge=0, le=10000)

    @model_validator(mode='after')
    def valid_name(self):
        if not self.name.strip():
            raise ValueError('Give the template a name.')
        return self


def assign_wells(plate: Plate, template: ReadingTemplate, sample_index: int):
    if sample_index < 0:
        raise ValueError('Sample index cannot be negative.')
    rows = plate.rows if template.row_direction == 'top' else plate.rows[::-1]
    columns = plate.columns if template.column_direction == 'left' else plate.columns[::-1]
    groups = [columns[i:i + template.replicates] for i in range(0, len(columns), template.replicates)]
    slots = [(row, group) for group in groups if len(group) == template.replicates for row in rows]
    usable = slots[template.start_slot:]
    samples_per_plate = len(usable) // template.flask_count
    if not samples_per_plate:
        raise ValueError('This grid cannot fit the selected flasks and replicates. Change the template or use a larger plate.')
    plate_number, sample_on_plate = divmod(sample_index, samples_per_plate)
    start = sample_on_plate * template.flask_count
    readings = {(well.row, well.column): well for well in plate.wells}
    result = []
    for flask_offset, (row, group) in enumerate(usable[start:start + template.flask_count]):
        for replicate, column in enumerate(group, 1):
            well = readings[row, column]
            result.append(dict(flask=f'FL-{template.flask_start + flask_offset:02d}', replicate=replicate,
                               row=row, column=column, value=well.value, status=well.status,
                               plate_number=plate_number + 1))
    return result


def trajectories(captures, started_at):
    start = datetime.fromisoformat(started_at)
    points = []
    for capture in captures:
        if capture['status'] != 'reviewed':
            continue
        elapsed = (datetime.fromisoformat(capture['captured_at']) - start).total_seconds() / 3600
        by_flask = defaultdict(list)
        for item in capture['readings']:
            by_flask[item['flask']].append(item)
        for flask, readings in by_flask.items():
            values = [item['value'] for item in readings if item['value'] is not None]
            complete = len(values) == len(readings)
            points.append(dict(flask=flask, captured_at=capture['captured_at'], hours=elapsed,
                               mean=mean(values) if complete else None,
                               sd=stdev(values) if complete and len(values) > 1 else None,
                               count=len(values), expected=len(readings), capture_id=capture['id']))
    return sorted(points, key=lambda item: (item['flask'], item['hours'], item['capture_id']))


def experiment_workbook(experiment, captures):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Readings'
    sheet.append(['Experiment', 'Captured at (UTC)', 'Elapsed hours', 'Photo', 'Template', 'Sample',
                  'Plate', 'Flask', 'Replicate', 'Row', 'Column', 'OD', 'Review status', 'Reading status'])
    start = datetime.fromisoformat(experiment['started_at'])
    for capture in captures:
        for reading in capture.get('readings', []):
            sheet.append([experiment['name'], capture['captured_at'],
                          (datetime.fromisoformat(capture['captured_at']) - start).total_seconds() / 3600,
                          capture['filename'], capture['template']['name'], capture['sample_index'] + 1,
                          reading['plate_number'], reading['flask'], reading['replicate'], reading['row'],
                          reading['column'], reading['value'], capture['status'], reading['status']])
    summary = workbook.create_sheet('Trajectories')
    summary.append(['Flask', 'Captured at (UTC)', 'Elapsed hours', 'Mean OD', 'Sample SD', 'Valid replicates', 'Expected replicates'])
    points = trajectories(captures, experiment['started_at'])
    for point in points:
        summary.append([point['flask'], point['captured_at'], point['hours'], point['mean'], point['sd'], point['count'], point['expected']])
    if points:
        chart = ScatterChart()
        chart.title = 'OD over time · reviewed readings'
        chart.scatterStyle = 'lineMarker'
        chart.x_axis.title = 'Hours from experiment start'
        chart.y_axis.title = 'Mean OD'
        chart.display_blanks = 'gap'
        for flask in sorted({point['flask'] for point in points}):
            indices = [i + 2 for i, point in enumerate(points) if point['flask'] == flask]
            series = Series(Reference(summary, min_col=4, min_row=min(indices), max_row=max(indices)),
                            Reference(summary, min_col=3, min_row=min(indices), max_row=max(indices)), title=flask)
            series.marker.symbol = 'circle'
            chart.series.append(series)
        summary.add_chart(chart, 'I2')
    photos = workbook.create_sheet('Photos')
    photos.append(['Photo ID', 'Filename', 'Captured at (UTC)', 'Original capture time', 'Template', 'Sample', 'Status', 'Error'])
    for capture in captures:
        photos.append([capture['id'], capture['filename'], capture['captured_at'], capture['original_captured_at'],
                       capture['template']['name'], capture['sample_index'] + 1, capture['status'], capture.get('error')])
    for i, capture in enumerate(captures, 1):
        if capture.get('plate'):
            plate = Plate.model_validate(capture['plate'])
            values = {(well.row, well.column): well.value for well in plate.wells}
            raw = workbook.create_sheet(f'Photo {i} raw')
            write_sheet(raw, [[values[row, col] for col in plate.columns] for row in plate.rows],
                        {(well.row, well.column) for well in plate.wells if well.status == 'unclear'}, plate.rows, plate.columns)
    for tab in workbook:
        tab.freeze_panes = 'B2'
        tab.auto_filter.ref = tab.dimensions
        for row in tab:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        for column in tab.columns:
            tab.column_dimensions[column[0].column_letter].width = min(32, max(14, len(str(column[0].value)) + 2))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
