import json
import os
import tempfile
import unittest
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from unittest.mock import MagicMock, patch
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

import server
from extraction import Plate, Well
from protocol import ReadingTemplate, assign_wells, experiment_workbook, trajectories
from sheets_sync import write_spreadsheet


def plate():
    return Plate(rows=list('ABCDEFGHI'), columns=[str(i) for i in range(1, 7)], wells=[
        Well(row=row, column=str(col), value=(r + 1) / 10 + col / 1000, status='readable')
        for r, row in enumerate('ABCDEFGHI') for col in range(1, 7)
    ])


class LabTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.patches = [patch.object(server, 'DATA', Path(self.directory.name)), patch.dict(os.environ, {'DISABLE_WORKER': '1', 'GEMINI_API_KEY': 'fake-key'})]
        for item in self.patches:
            item.start()
        self.client = TestClient(server.app)
        self.client.__enter__()
        self.client.headers['X-Plate-Request'] = '1'
        self.experiment = self.client.post('/api/experiments', json={'name': 'EXP-01', 'started_at': '2026-10-02T09:00:00+05:30'}).json()
        self.path = '/api/experiments/' + self.experiment['id']
        self.template = self.client.post(self.path + '/templates', json={'name': 'First eight', 'flask_count': 8}).json()

    def tearDown(self):
        self.client.__exit__(None, None, None)
        for item in reversed(self.patches):
            item.stop()
        self.directory.cleanup()

    def upload(self, capture_id=None):
        image = BytesIO()
        Image.new('RGB', (30, 30), 'white').save(image, 'JPEG')
        response = self.client.post(self.path + '/photos', data={'id': capture_id or str(uuid4()), 'template_id': self.template['id'], 'captured_at': '2026-10-02T10:00:00+05:30'}, files={'photo': ('capture.jpg', image.getvalue(), 'image/jpeg')})
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def extract(self, capture):
        self.assertEqual(self.client.post('/api/photos/' + capture['id'] + '/process').status_code, 200)
        with patch('server.extract_plate', return_value=plate()) as call:
            self.assertTrue(server.process_one())
            call.assert_called_once()
        return self.client.get(self.path).json()['captures'][0]

    def test_photo_idempotency_review_export_and_history(self):
        capture = self.upload()
        self.assertEqual(self.upload(capture['id'])['sample_index'], 0)
        capture = self.extract(capture)
        self.assertEqual(capture['status'], 'ready')
        self.assertEqual(capture['readings'][0]['row'], 'I')
        self.assertEqual(capture['readings'][0]['column'], '1')
        data = dict(revision=capture['revision'], captured_at=capture['captured_at'], sample_index=0, readings=capture['readings'])
        data['readings'][0]['value'] = 0.321
        response = self.client.put('/api/photos/' + capture['id'] + '/review', json=data)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(self.client.put('/api/photos/' + capture['id'] + '/review', json=data).status_code, 409)
        experiment = self.client.get(self.path).json()
        self.assertEqual(len(experiment['trajectories']), 8)
        self.assertEqual(experiment['trajectories'][0]['hours'], 1)
        saved = experiment['captures'][0]
        self.assertEqual(saved['readings'][0]['status'], 'corrected')
        self.assertNotEqual(saved['plate']['wells'][-6]['value'], 0.321)
        workbook = load_workbook(BytesIO(self.client.get(self.path + '/export').content))
        self.assertIn('Photo 1 raw', workbook.sheetnames)
        self.assertEqual(len(workbook['Trajectories']._charts), 1)
        self.assertEqual(workbook['Readings']['L2'].value, .321)
        with server.database() as db:
            self.assertEqual(db.execute('SELECT COUNT(*) FROM edits').fetchone()[0], 1)

    def test_offline_timestamp_failure_retry_and_restart(self):
        capture = self.upload()
        self.assertEqual(capture['captured_at'], '2026-10-02T04:30:00+00:00')
        self.client.post('/api/photos/' + capture['id'] + '/process')
        with patch('server.extract_plate', side_effect=RuntimeError('network disconnected')):
            server.process_one()
        self.assertEqual(self.client.get(self.path).json()['captures'][0]['status'], 'failed')
        self.extract(capture)
        server.init_db()
        self.assertEqual(self.client.get(self.path).json()['captures'][0]['status'], 'ready')

    def test_templates_remap_and_missing_replicates(self):
        capture = self.extract(self.upload())
        other = self.client.post(self.path + '/templates', json={'name': 'Flasks 09–16', 'flask_start': 9, 'flask_count': 8, 'row_direction': 'top'}).json()
        response = self.client.post('/api/photos/' + capture['id'] + '/remap', json={'revision': capture['revision'], 'template_id': other['id'], 'sample_index': 1})
        self.assertEqual(response.status_code, 200, response.text)
        capture = self.client.get(self.path).json()['captures'][0]
        self.assertEqual(capture['readings'][0]['flask'], 'FL-09')
        self.assertEqual(capture['readings'][0]['row'], 'I')
        capture['readings'][0]['value'] = None
        response = self.client.put('/api/photos/' + capture['id'] + '/review', json={'revision': capture['revision'], 'captured_at': capture['captured_at'], 'sample_index': capture['sample_index'], 'readings': capture['readings']})
        self.assertEqual(response.status_code, 200, response.text)
        points = self.client.get(self.path).json()['trajectories']
        self.assertIsNone(points[0]['mean'])
        self.assertEqual(points[0]['count'], 1)

    def test_request_validation(self):
        response = self.client.post(self.path + '/templates', json={'name':'Invalid','replicates':0})
        self.assertEqual(response.status_code, 422)
        self.client.headers.pop('X-Plate-Request')
        self.assertEqual(self.client.post(self.path + '/sync').status_code, 403)

    def test_google_sheet_sync_contract_and_recovery(self):
        capture = self.extract(self.upload())
        data = dict(revision=capture['revision'], captured_at=capture['captured_at'], sample_index=0, readings=capture['readings'])
        self.client.put('/api/photos/' + capture['id'] + '/review', json=data)
        self.client.put(self.path, json={'name': 'EXP-01', 'started_at': self.experiment['started_at'], 'sheet_id': 'example-sheet-id-1234567890'})
        experiment = self.client.get(self.path).json()
        session = MagicMock()
        session.get.return_value.json.return_value = {'sheets': []}
        write_spreadsheet(session, experiment, experiment['captures'])
        requests = session.post.call_args.kwargs['json']['requests']
        self.assertEqual(sum('addSheet' in item for item in requests), 3)
        self.assertEqual(sum('addChart' in item for item in requests), 8)
        self.assertEqual(sum('updateCells' in item for item in requests), 3)
        with patch('sheets_sync.sync_experiment', side_effect=RuntimeError('no access')):
            server.sync_one()
        self.assertEqual(self.client.get(self.path).json()['sync_status'], 'failed')
        self.client.post(self.path + '/sync')
        with patch('sheets_sync.sync_experiment'):
            self.assertTrue(server.sync_one())
        self.assertEqual(self.client.get(self.path).json()['sync_status'], 'synced')


class ProtocolTests(unittest.TestCase):
    def test_sequence_rollover_and_nonstandard_grid(self):
        template = ReadingTemplate(name='16 flasks', flask_count=16)
        first = assign_wells(plate(), template, 0)
        self.assertEqual((first[0]['row'], first[0]['column']), ('I', '1'))
        self.assertEqual((first[18]['row'], first[18]['column']), ('I', '3'))
        self.assertEqual(assign_wells(plate(), template, 1)[0]['plate_number'], 2)
        other = template.model_copy(update={'replicates': 3, 'flask_count': 2, 'column_direction': 'right', 'start_slot': 1})
        self.assertEqual([item['column'] for item in assign_wells(plate(), other, 0)[:3]], ['6','5','4'])
        with self.assertRaises(ValueError):
            assign_wells(plate(), ReadingTemplate(name='too many', flask_count=100), 0)


if __name__ == '__main__':
    unittest.main()
